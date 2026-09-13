#!/usr/bin/env python3
"""Названия сообществ построчно -> ссылки vk.ru, Excel и список слагов.

Обычный интернет-поиск без токенов и браузера. Зависимость: openpyxl.
"""

import argparse
from collections import Counter
from difflib import SequenceMatcher
import fcntl
import html
from html.parser import HTMLParser
import json
import os
from pathlib import Path
import re
import sys
import time
import unicodedata
from urllib.error import HTTPError, URLError
from urllib.parse import parse_qs, unquote, urlencode, urlparse
from urllib.request import Request, urlopen


ROOT = Path(__file__).resolve().parent
DEFAULT_SOURCE = ROOT / "vk.fastfixsite.ru/public_html/sourse.html"
DEFAULT_OUTPUT = ROOT / "output/vk-groups-web"
USER_AGENT = ("Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/136.0.0.0 Safari/537.36")
# Публичный поисковый виджет Blackle, работающий на Google CSE.
# Это идентификатор поисковика, а не пользовательский API-ключ.
GOOGLE_CX = "partner-pub-8993703457585266:4862972284"
STATUSES = {
    "found": "Совпадение названия в поисковой выдаче",
    "ambiguous": "Несколько одинаковых названий — проверить",
    "review": "Точного совпадения нет — проверить варианты",
    "missing": "В выдаче нет ссылок на страницы VK",
    "pending": "Ещё не обработано",
}


def clean_name(value):
    value = html.unescape(value.strip())
    if len(value) >= 2 and value.startswith('"') and value.endswith('"'):
        value = value[1:-1].replace('""', '"')
    return " ".join(value.split())


def normalize(value):
    return unicodedata.normalize("NFKC", clean_name(value)).casefold()


def read_names(path):
    # Расширение .html несущественно: в исходном файле одна строка = название.
    return [
        (number, line.strip(), clean_name(line))
        for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1)
        if clean_name(line)
    ]


class SearchError(Exception):
    pass


class TemporarySearchError(SearchError):
    """Временный сбой соединения; остальные названия можно обработать."""


class ResultParser(HTMLParser):
    """Только заголовки результатов, без меню, рекламных блоков и сниппетов."""

    def __init__(self, engine):
        super().__init__(convert_charrefs=True)
        self.engine = engine
        self.results = []
        self.anchor = None
        self.parts = []
        self.heading = False

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == "a":
            self.anchor = attrs.get("href", "")
            self.parts = []
            self.heading = self.engine == "duckduckgo" and "result__a" in attrs.get("class", "").split()
        elif tag == "h3" and self.anchor is not None and self.engine == "google":
            self.heading = True
            self.parts = []

    def handle_data(self, data):
        if self.anchor is not None and self.heading:
            self.parts.append(data)

    def handle_endtag(self, tag):
        if self.engine == "google" and tag == "h3" and self.heading:
            self.finish_result()
            self.heading = False
        if tag == "a":
            if self.heading:
                self.finish_result()
            self.anchor = None
            self.parts = []
            self.heading = False

    def finish_result(self):
        title = " ".join("".join(self.parts).split())
        if self.anchor and title:
            self.results.append({"href": self.anchor, "title": title})


def unwrap_url(value):
    value = html.unescape(value)
    if value.startswith("//"):
        value = "https:" + value
    parsed = urlparse(value)
    params = parse_qs(parsed.query)
    host = (parsed.hostname or "").lower()
    if host in ("duckduckgo.com", "html.duckduckgo.com") and params.get("uddg"):
        return params["uddg"][0]
    if parsed.path == "/url" and (not host or host in ("www.google.com", "google.com")):
        return (params.get("q") or params.get("url") or [value])[0]
    return value


def candidate_from_result(result):
    source = unwrap_url(result["href"])
    parsed = urlparse(source)
    host = (parsed.hostname or "").lower()
    path = unquote(parsed.path).strip("/")
    if parsed.scheme not in ("http", "https") or parsed.username or parsed.password:
        return None
    if host not in ("vk.com", "www.vk.com", "m.vk.com", "vk.ru", "www.vk.ru", "m.vk.ru"):
        return None
    # Не превращаем ссылки на посты, видео, поиск, вход и личные ID в группы.
    reserved = {
        "feed", "im", "login", "join", "search", "away", "away.php", "share", "share.php",
        "video", "videos", "clips", "photo", "photos", "audio", "audios", "friends", "groups",
        "public", "club", "wall", "market", "catalog", "apps", "app", "settings", "help",
        "about", "support", "restore", "terms", "privacy", "legal", "al_video.php", "al_wall.php",
    }
    if not re.fullmatch(r"[A-Za-z0-9_.]+", path) or path.lower() in reserved:
        return None
    if re.match(r"^(?:id\d|event\d|wall-?\d|video-?\d|clip-?\d|photo-?\d|album-?\d|topic-?\d|app\d)", path, re.I):
        return None
    if any(key in parse_qs(parsed.query) for key in ("w", "z")):
        return None
    return {"name": result["title"], "slug": path,
            "url": f"https://vk.ru/{path}", "source_url": source}


class WebSearch:
    def __init__(self, engine="google", delay=5):
        self.engine = engine
        self.delay = delay
        self.last_request = 0
        self.google_config = None
        self.google_config_time = 0

    def fetch(self, url):
        request = Request(url, headers={
            "User-Agent": USER_AGENT, "Accept-Language": "ru,en;q=0.8",
            "Referer": "https://cse.google.com/" if self.engine == "google" else "https://duckduckgo.com/",
        })
        error_message = "Поисковик не ответил."
        for attempt in range(3):
            time.sleep(max(0, self.delay - (time.monotonic() - self.last_request)))
            self.last_request = time.monotonic()
            try:
                with urlopen(request, timeout=20) as response:
                    body = response.read().decode("utf-8", errors="replace")
                    final_url = response.geturl()
                    http_status = response.status
            except HTTPError as exc:
                error_message = f"{self.engine}: HTTP {exc.code}. Повторите запуск позже."
                if exc.code < 500:
                    raise SearchError(error_message) from None
            except (URLError, TimeoutError, OSError, ValueError):
                error_message = f"Нет соединения с {self.engine}. Проверьте интернет или смените --engine."
            else:
                lower = body.lower()
                blockers = ("anomaly-modal", "anomaly.js", "g-recaptcha", "unusual traffic",
                            "докажите, что вы не робот", "enablejs")
                if http_status in (202, 429) or any(marker in lower for marker in blockers) or "/sorry/" in final_url or "consent.google" in final_url:
                    raise SearchError(f"{self.engine}: капча, ограничение или требуется JavaScript. "
                                      "Текущая строка останется необработанной.")
                return body
            if attempt < 2:
                time.sleep(3 * (attempt + 1))
        raise TemporarySearchError(error_message)

    def search(self, query):
        if self.engine == "google":
            return self.google_search(query)
        body = self.fetch("https://html.duckduckgo.com/html/?" + urlencode({"q": query}))
        parser = ResultParser("duckduckgo")
        parser.feed(body)
        if parser.results:
            return parser.results
        if any(marker in body.lower() for marker in ("no-results", "no results found")):
            return []
        raise SearchError("duckduckgo: не удалось распознать выдачу. Строка останется необработанной.")

    def google_search(self, query):
        if self.google_config is None or time.monotonic() - self.google_config_time >= 1800:
            script = self.fetch("https://www.google.ru/cse/cse.js?" + urlencode({"cx": GOOGLE_CX}))
            try:
                # Параметры публичного виджета доступны без аккаунта; JS не выполняется.
                begin, end = script.rfind("({"), script.rfind("});")
                config = json.loads(script[begin + 1:end + 1])
                if not isinstance(config, dict) or not config.get("cse_token"):
                    raise ValueError
            except ValueError:
                raise SearchError("Google: не удалось прочитать настройки публичного поиска.") from None
            self.google_config = config
            self.google_config_time = time.monotonic()
        config = self.google_config
        params = {
            "cx": GOOGLE_CX, "q": query, "hl": "ru", "safe": "off",
            "num": 20, "rsz": "filtered_cse", "callback": "_",
            "rurl": "", "searchtype": "", "cselibv": config.get("cselibVersion", ""),
            "cse_tok": config["cse_token"],
        }
        if config.get("exp"):
            params["exp"] = ",".join(config["exp"])
        body = self.fetch("https://cse.google.com/cse/element/v1?" + urlencode(params))
        try:
            response = json.loads(body[body.find("{"):body.rfind("}") + 1])
        except ValueError:
            raise SearchError("Google: не удалось прочитать результаты поиска.") from None
        if not isinstance(response, dict):
            raise SearchError("Google: неожиданный формат ответа.")
        if response.get("error"):
            # Не печатаем ответ целиком: в нём могут повторяться временные параметры виджета.
            code = response["error"].get("code", "неизвестная")
            raise SearchError(f"Google: ошибка {code}. Текущая строка останется необработанной.")
        items = response.get("results")
        if items is None and "cursor" in response:
            return []
        if not isinstance(items, list):
            raise SearchError("Google: в ответе нет списка результатов.")
        return [{"href": item["unescapedUrl"], "title": html.unescape(item.get("titleNoFormatting", ""))}
                for item in items if isinstance(item, dict) and item.get("unescapedUrl")]


def title_key(value):
    value = re.sub(r"\s*[|—–-]\s*(?:VK|ВКонтакте)\s*$", "", value, flags=re.I)
    return " ".join(re.findall(r"\w+", normalize(value)))


def title_matches(name, title):
    target, found = title_key(name), title_key(title)
    # В заголовке выдачи после названия часто идёт начало описания страницы.
    return bool(target) and (found == target or (
        len(target.split()) >= 3 and found.startswith(target + " ")))


def find_group(name, search):
    quoted = name.replace('"', ' ')
    queries = [f'"{quoted}" (site:vk.com OR site:vk.ru)', f"{name} вконтакте"]
    prefix = re.split(r"[|/]", name, maxsplit=1)[0].strip()
    if prefix != name and (len(prefix.split()) >= 2 or re.fullmatch(r"[A-Za-z][\w.]{5,}", prefix)):
        queries.append(f'"{prefix}" вконтакте')
    candidates = {}
    for query in queries:
        for item in search.search(query):
            group = candidate_from_result(item)
            if group:
                key = group["slug"].casefold()
                if key not in candidates or title_matches(name, group["name"]):
                    candidates[key] = group
        exact = [g for g in candidates.values() if title_matches(name, g["name"])]
        if exact:
            return {"status": "found" if len(exact) == 1 else "ambiguous",
                    "candidates": exact, "query": query, "engine": search.engine}
    ranked = sorted(candidates.values(), key=lambda g: SequenceMatcher(
        None, title_key(name), title_key(g["name"]), autojunk=False).ratio(), reverse=True)
    return {"status": "review" if ranked else "missing", "candidates": ranked[:5],
            "query": queries[-1], "engine": search.engine}


def save_progress(path, records):
    temporary = path.with_suffix(".tmp")
    temporary.write_text(json.dumps({"version": 2, "records": records},
                                    ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(path)


def load_progress(path):
    if not path.exists():
        return {}
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or data.get("version") != 2 or not isinstance(data.get("records"), dict):
        raise ValueError("Неизвестный формат progress.json; выберите другую папку --output.")
    records = data["records"]
    for record in records.values():
        if not isinstance(record, dict) or record.get("status") not in STATUSES:
            raise ValueError("Некорректная запись в progress.json.")
        groups = record.get("candidates")
        if not isinstance(groups, list) or any(
            not isinstance(g, dict) or not all(k in g for k in ("name", "slug", "url"))
            for g in groups
        ) or (record["status"] == "found" and len(groups) != 1):
            raise ValueError("Некорректные кандидаты в progress.json.")
    return records


def export_results(rows, records, output):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    links = workbook.active
    links.title = "Ссылки"
    slugs = workbook.create_sheet("Слаги")
    review = workbook.create_sheet("Проверить")

    def append_text(sheet, values):
        # Названия вроде '=...' должны оставаться текстом, а не формулами Excel.
        sheet.append([str(value) for value in values])
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"

    append_text(links, ["Ссылка — название", "Ссылка", "Заголовок в поиске", "Слаг",
                        "Строка файла", "Исходная строка", "Статус", "Поиск в Google"])
    append_text(slugs, ["Слаг"])
    append_text(review, ["Строка файла", "Исходная строка", "Статус",
                         "Вариант: ссылка", "Вариант: заголовок", "Вариант: слаг", "Поиск в Google"])
    link_lines, slug_lines, seen_slugs = [], [], set()
    for number, original, name in rows:
        record = records.get(name, {"status": "pending", "candidates": []})
        status = STATUSES[record["status"]]
        google_url = "https://www.google.com/search?" + urlencode({"q": f"{name} вконтакте"})
        if record["status"] == "found":
            group = record["candidates"][0]
            label = f'{group["url"]} - {name}'
            append_text(links, [label, group["url"], group["name"], group["slug"],
                                number, original, status, google_url])
            links.cell(links.max_row, 2).hyperlink = group["url"]
            links.cell(links.max_row, 2).style = "Hyperlink"
            link_lines.append(label)
            if group["slug"].casefold() not in seen_slugs:
                seen_slugs.add(group["slug"].casefold())
                slug_lines.append(group["slug"])
                append_text(slugs, [group["slug"]])
        else:
            append_text(links, ["", "", "", "", number, original, status, google_url])
            for group in record["candidates"] or [{}]:
                append_text(review, [number, original, status, group.get("url", ""),
                                     group.get("name", ""), group.get("slug", ""), google_url])
                if group.get("url"):
                    review.cell(review.max_row, 4).hyperlink = group["url"]
                    review.cell(review.max_row, 4).style = "Hyperlink"
                review.cell(review.max_row, 7).hyperlink = google_url
                review.cell(review.max_row, 7).style = "Hyperlink"
        links.cell(links.max_row, 8).hyperlink = google_url
        links.cell(links.max_row, 8).style = "Hyperlink"

    for sheet, widths in ((links, [85, 38, 60, 28, 16, 65, 50, 30]),
                          (slugs, [35]), (review, [16, 65, 50, 38, 65, 28, 30])):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0077FF")
        for column, width in enumerate(widths, 1):
            sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    temporary = output / "vk_groups.tmp.xlsx"
    workbook.save(temporary)
    temporary.replace(output / "vk_groups.xlsx")
    for filename, lines in (("links.txt", link_lines), ("slugs.txt", slug_lines)):
        temporary = output / (filename + ".tmp")
        temporary.write_text("\n".join(lines) + ("\n" if lines else ""), encoding="utf-8")
        temporary.replace(output / filename)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", type=Path, default=DEFAULT_SOURCE,
                        help="UTF-8 файл: одно название на строку")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT,
                        help="Папка для Excel, txt и прогресса")
    parser.add_argument("--limit", type=int, help="Обработать только первые N строк")
    parser.add_argument("--engine", choices=("google", "duckduckgo"), default="google",
                        help="Поисковик; по умолчанию публичный поиск Google CSE")
    parser.add_argument("--delay", type=float, default=5,
                        help="Пауза между запросами, минимум 3 секунды")
    parser.add_argument("--export-only", action="store_true",
                        help="Пересобрать файлы из прогресса без интернета и токена")
    parser.add_argument("--retry-unresolved", action="store_true",
                        help="Повторить поиск для ненайденных и неоднозначных названий")
    args = parser.parse_args(argv)
    if args.limit is not None and args.limit < 1:
        parser.error("--limit должен быть больше нуля")
    if not 3 <= args.delay <= 60:
        parser.error("--delay должен быть от 3 до 60 секунд")
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("Установите зависимость: python3 -m pip install openpyxl", file=sys.stderr)
        return 1
    worker_lock = None
    try:
        rows = read_names(args.source.expanduser())
        if args.limit:
            rows = rows[:args.limit]
        if not rows:
            raise ValueError("В исходном файле нет названий.")
        output = args.output.expanduser().resolve()
        protected = {output / f for f in ("vk_groups.xlsx", "links.txt", "slugs.txt",
                                          "progress.json", "progress.tmp")}
        if args.source.expanduser().resolve() in protected:
            raise ValueError("Исходный файл не должен совпадать с выходным.")
        output.mkdir(parents=True, exist_ok=True)
        worker_lock = (output / ".worker.lock").open("a+")
        try:
            fcntl.flock(worker_lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError:
            print("Сбор или выгрузка уже выполняется в этой папке. Второй процесс не запущен.", file=sys.stderr)
            return 75
        worker_lock.seek(0)
        worker_lock.truncate()
        worker_lock.write(str(os.getpid()))
        worker_lock.flush()
        progress = output / "progress.json"
        records = load_progress(progress)
        names = list(dict.fromkeys(name for _, _, name in rows))
        pending = [name for name in names if name not in records or records[name]["status"] == "pending" or (
            args.retry_unresolved and records[name]["status"] != "found")]
        print(f"Строк: {len(rows)}; уникальных названий: {len(names)}; осталось: {len(pending)}.")
        exit_code = 0
        try:
            if pending and not args.export_only:
                search = WebSearch(args.engine, args.delay)
                queue = pending
                consecutive_failures = 0
                for pass_number in range(2):
                    deferred = []
                    for index, name in enumerate(queue, 1):
                        label = "Ищу" if pass_number == 0 else "Повторяю после сетевого сбоя"
                        print(f"[{index}/{len(queue)}] {label}: {name}", flush=True)
                        try:
                            record = find_group(name, search)
                        except TemporarySearchError as exc:
                            deferred.append(name)
                            consecutive_failures += 1
                            print(f"Временный сбой: {name}. {exc}", file=sys.stderr, flush=True)
                            if consecutive_failures >= 3:
                                raise SearchError("Три сетевых сбоя подряд. Остановлено; оставшиеся строки можно продолжить повторным запуском.") from None
                            continue
                        consecutive_failures = 0
                        records[name] = record
                        save_progress(progress, records)
                        print(f"[{index}/{len(queue)}] {STATUSES[record['status']]}: {name}", flush=True)
                        if index % 25 == 0:
                            export_results(rows, records, output)
                    if not deferred:
                        break
                    queue = deferred
                else:
                    print(f"Не удалось повторно обработать строк из-за сетевых сбоев: {len(queue)}.", file=sys.stderr)
                    exit_code = 1
        except SearchError as exc:
            print(f"Поиск остановлен: {exc}", file=sys.stderr)
            exit_code = 1
        except KeyboardInterrupt:
            print("\nОстановлено. Повторный запуск продолжит с сохранённого места.")
            exit_code = 130
        finally:
            export_results(rows, records, output)
        counts = Counter(records.get(name, {}).get("status", "pending") for _, _, name in rows)
        print(f"Ссылок: {counts['found']}; проверить: {counts['ambiguous'] + counts['review']}; "
              f"не найдено: {counts['missing']}; не обработано: {counts['pending']}.")
        print(f"Файлы: {output}")
        return exit_code
    except (OSError, ValueError) as exc:
        print(f"Ошибка: {exc}", file=sys.stderr)
        return 1
    finally:
        if worker_lock is not None:
            worker_lock.close()


if __name__ == "__main__":
    sys.exit(main())
